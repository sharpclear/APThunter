"""Match newly registered domains against APT-style registration templates.

Templates are read from the workbook produced by the APT domain template
screening step. A date range such as 20260301-20260302 selects daily NRD zip
files under dataset/NRD/YYYY-MM/YYYY-MM-DD-domain.zip and writes candidate
matches to outputs/apt_template_nrd_matcher/<range>/.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import time
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Iterator

try:
    import openpyxl

    HAS_OPENPYXL = True
except Exception:
    openpyxl = None
    HAS_OPENPYXL = False

try:
    import template_batch_detector as domain_utils
except ImportError:
    from . import template_batch_detector as domain_utils  # type: ignore

import pandas as pd


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATES_PATH = SCRIPT_DIR / "dataset" / "APTdomains" / "疑似模板化注册域名_筛选结果.xlsx"
DEFAULT_NRD_ROOT = SCRIPT_DIR / "dataset" / "NRD"
DEFAULT_OUTPUT_ROOT = SCRIPT_DIR / "outputs"
OUTPUT_JOB_NAME = "apt_template_nrd_matcher"
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "latin-1")
DEFAULT_SCORE_THRESHOLD = 0.90


def _common_multi_suffixes() -> set[str]:
    return set(
        getattr(
            domain_utils,
            "COMMON_MULTI_SUFFIXES",
            getattr(domain_utils, "_KNOWN_MULTI_PART_SUFFIXES", set()),
        )
    )


def _parse_domain_value(domain: str) -> Any:
    parser = getattr(domain_utils, "parse_normalized_domain", None)
    if parser is None:
        parser = getattr(domain_utils, "parse_domain")
    return parser(domain)


@dataclass(frozen=True)
class TemplateSpec:
    template: str
    sld_template: str
    suffix_template: str
    reason: str
    root_count: int
    row_count: int
    sld_regex: re.Pattern[str]
    wildcard_count: int
    fixed_token_count: int


@dataclass(frozen=True)
class TemplateMatch:
    template: str
    reason: str
    variables: dict[str, str]
    root_count: int
    row_count: int


@dataclass(frozen=True)
class DailyNrdFile:
    day: date
    path: Path


def parse_date_range(value: str) -> list[date]:
    """Parse inclusive YYYYMMDD-YYYYMMDD date ranges."""
    text = value.strip()
    if not re.fullmatch(r"\d{8}-\d{8}", text):
        raise ValueError("date range must use YYYYMMDD-YYYYMMDD, for example 20260301-20260302")
    start_text, end_text = text.split("-", 1)
    start = datetime.strptime(start_text, "%Y%m%d").date()
    end = datetime.strptime(end_text, "%Y%m%d").date()
    if start > end:
        raise ValueError(f"date range start is after end: {value}")
    days: list[date] = []
    current = start
    while current <= end:
        days.append(current)
        current += timedelta(days=1)
    return days


def date_range_slug(value: str) -> str:
    parse_date_range(value)
    return value.strip()


def load_templates(path: str | Path) -> list[TemplateSpec]:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required to read template xlsx files")
    template_path = Path(path)
    if not template_path.exists():
        raise FileNotFoundError(f"template workbook does not exist: {template_path}")

    workbook = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
    try:
        rows = _read_template_rows(workbook)
    finally:
        workbook.close()

    specs: list[TemplateSpec] = []
    seen: set[str] = set()
    for row in rows:
        template = str(row.get("模板") or "").strip()
        if not template or template in seen:
            continue
        seen.add(template)
        specs.append(
            compile_template(
                template,
                reason=str(row.get("判定依据") or ""),
                root_count=_safe_int(row.get("命中注册域名数")),
                row_count=_safe_int(row.get("命中原始域名行数")),
            )
        )
    if not specs:
        raise ValueError(f"no templates found in workbook: {template_path}")
    return specs


def _read_template_rows(workbook: Any) -> list[dict[str, Any]]:
    if "模板汇总" in workbook.sheetnames:
        return _rows_from_sheet(workbook["模板汇总"])
    if "筛选结果" in workbook.sheetnames:
        rows = _rows_from_sheet(workbook["筛选结果"])
        collapsed: dict[str, dict[str, Any]] = {}
        for row in rows:
            template = str(row.get("模板") or "").strip()
            if not template:
                continue
            current = collapsed.setdefault(
                template,
                {
                    "模板": template,
                    "命中注册域名数": 0,
                    "命中原始域名行数": 0,
                    "判定依据": row.get("判定依据") or "",
                },
            )
            current["命中原始域名行数"] = int(current["命中原始域名行数"]) + 1
        return list(collapsed.values())
    first_sheet = workbook[workbook.sheetnames[0]]
    return _rows_from_sheet(first_sheet)


def _rows_from_sheet(sheet: Any) -> list[dict[str, Any]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = next(iterator)
    except StopIteration:
        return []
    header_names = [str(header).strip() if header is not None else "" for header in headers]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if not values or all(value in (None, "") for value in values):
            continue
        rows.append({header_names[index]: value for index, value in enumerate(values) if index < len(header_names)})
    return rows


def compile_template(template: str, *, reason: str = "", root_count: int = 0, row_count: int = 0) -> TemplateSpec:
    sld_template, suffix_template = split_template(template)
    regex_parts: list[str] = []
    wildcard_index = 0
    fixed_token_count = 0

    for part in sld_template.split("-"):
        if part == "{x}":
            wildcard_index += 1
            regex_parts.append(f"(?P<x{wildcard_index}>[a-z0-9]+)")
        else:
            fixed_token_count += 1
            regex_parts.append(re.escape(part))
    sld_regex = re.compile(r"^" + "-".join(regex_parts) + r"$", re.IGNORECASE)
    return TemplateSpec(
        template=template,
        sld_template=sld_template,
        suffix_template=suffix_template,
        reason=reason,
        root_count=root_count,
        row_count=row_count,
        sld_regex=sld_regex,
        wildcard_count=wildcard_index,
        fixed_token_count=fixed_token_count,
    )


def split_template(template: str) -> tuple[str, str]:
    text = template.strip().lower()
    if text.endswith(".{tld}"):
        return text[: -len(".{tld}")], "{tld}"
    labels = text.split(".")
    if len(labels) >= 3 and ".".join(labels[-2:]) in _common_multi_suffixes():
        return ".".join(labels[:-2]), ".".join(labels[-2:])
    if len(labels) >= 2:
        return ".".join(labels[:-1]), labels[-1]
    raise ValueError(f"invalid template: {template}")


def match_domain(domain: str, templates: Iterable[TemplateSpec]) -> TemplateMatch | None:
    normalized = domain_utils.normalize_domain(domain)
    if not domain_utils.is_normalized_domain_like(normalized):
        return None
    try:
        parsed = _parse_domain_value(normalized)
    except Exception:
        return None

    matches: list[tuple[int, TemplateSpec, dict[str, str]]] = []
    for template in templates:
        variables = _match_template(parsed.sld, parsed.suffix, template)
        if variables is None:
            continue
        matches.append((_specificity_score(template), template, variables))
    if not matches:
        return None

    _, template, variables = max(matches, key=lambda item: item[0])
    return TemplateMatch(
        template=template.template,
        reason=template.reason,
        variables=variables,
        root_count=template.root_count,
        row_count=template.row_count,
    )


def _match_template(sld: str, suffix: str, template: TemplateSpec) -> dict[str, str] | None:
    if template.suffix_template != "{tld}" and suffix.lower() != template.suffix_template.lower():
        return None
    match = template.sld_regex.fullmatch(sld)
    if not match:
        return None
    variables = {key: value for key, value in match.groupdict().items() if value is not None}
    if template.suffix_template == "{tld}":
        variables["tld"] = suffix
    return variables


def _specificity_score(template: TemplateSpec) -> int:
    exact_sld_bonus = 10000 if template.wildcard_count == 0 else 0
    fixed_suffix_bonus = 1000 if template.suffix_template != "{tld}" else 0
    return exact_sld_bonus + fixed_suffix_bonus + template.fixed_token_count * 100 - template.wildcard_count


def locate_nrd_files(date_range: str, nrd_root: str | Path, *, allow_missing: bool = False) -> list[DailyNrdFile]:
    root = Path(nrd_root)
    files: list[DailyNrdFile] = []
    missing: list[Path] = []
    for day in parse_date_range(date_range):
        month = day.strftime("%Y-%m")
        dashed = day.strftime("%Y-%m-%d")
        path = root / month / f"{dashed}-domain.zip"
        if path.exists():
            files.append(DailyNrdFile(day=day, path=path))
        else:
            missing.append(path)
    if missing and not allow_missing:
        missing_text = "\n".join(str(path) for path in missing)
        raise FileNotFoundError(f"missing NRD files:\n{missing_text}")
    if not files:
        raise FileNotFoundError(f"no NRD files found for date range {date_range} under {root}")
    return files


def iter_nrd_zip_values(path: Path) -> Iterator[str]:
    with zipfile.ZipFile(path, "r") as archive:
        for name in sorted(archive.namelist()):
            if name.endswith("/"):
                continue
            suffix = Path(name).suffix.lower()
            if suffix not in {".txt", ".csv"}:
                continue
            text = decode_bytes(archive.read(name))
            if suffix == ".csv":
                for row in csv.reader(text.splitlines()):
                    if row:
                        yield row[0]
            else:
                for line in text.splitlines():
                    if line.strip():
                        yield line.strip()


def decode_bytes(raw: bytes) -> str:
    for encoding in TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def run_detection(
    *,
    date_range: str,
    templates_path: str | Path = DEFAULT_TEMPLATES_PATH,
    nrd_root: str | Path = DEFAULT_NRD_ROOT,
    output_root: str | Path = DEFAULT_OUTPUT_ROOT,
    allow_missing: bool = False,
) -> dict[str, Any]:
    started = time.time()
    slug = date_range_slug(date_range)
    templates = load_templates(templates_path)
    nrd_files = locate_nrd_files(slug, nrd_root, allow_missing=allow_missing)
    output_dir = Path(output_root) / OUTPUT_JOB_NAME / slug
    output_dir.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, Any]] = []
    total_input_rows = 0
    scanned_domains = 0
    invalid_rows = 0
    template_counts: Counter[str] = Counter()

    for daily_file in nrd_files:
        for raw_value in iter_nrd_zip_values(daily_file.path):
            total_input_rows += 1
            normalized = domain_utils.normalize_domain(raw_value)
            if not domain_utils.is_normalized_domain_like(normalized):
                invalid_rows += 1
                continue
            scanned_domains += 1
            parsed = _parse_domain_value(normalized)
            match = match_domain(normalized, templates)
            if match is None:
                continue
            registered_domain = f"{parsed.sld}.{parsed.suffix}" if parsed.suffix else parsed.sld
            template_counts[match.template] += 1
            rows.append(
                {
                    "date": daily_file.day.isoformat(),
                    "source_file": str(daily_file.path),
                    "domain": normalized,
                    "registered_domain": registered_domain,
                    "sld": parsed.sld,
                    "suffix": parsed.suffix,
                    "matched_template": match.template,
                    "template_reason": match.reason,
                    "template_root_count": match.root_count,
                    "template_row_count": match.row_count,
                    "variables_json": json.dumps(match.variables, ensure_ascii=False, sort_keys=True),
                }
            )

    candidates_path = output_dir / "apt_template_nrd_matches.csv"
    summary_path = output_dir / "apt_template_nrd_summary.json"
    write_matches(candidates_path, rows)

    summary: dict[str, Any] = {
        "date_range": slug,
        "templates_path": str(Path(templates_path).resolve()),
        "nrd_root": str(Path(nrd_root).resolve()),
        "output_dir": str(output_dir.resolve()),
        "output_matches_file": str(candidates_path.resolve()),
        "summary_file": str(summary_path.resolve()),
        "template_count": len(templates),
        "nrd_files": [str(item.path.resolve()) for item in nrd_files],
        "total_input_rows": total_input_rows,
        "scanned_domains": scanned_domains,
        "invalid_rows": invalid_rows,
        "matched_domains": len(rows),
        "template_match_counts": dict(sorted(template_counts.items())),
        "parameters": {"allow_missing": allow_missing},
        "run_seconds": round(time.time() - started, 3),
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def write_matches(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames = [
        "date",
        "source_file",
        "domain",
        "registered_domain",
        "sld",
        "suffix",
        "matched_template",
        "template_reason",
        "template_root_count",
        "template_row_count",
        "variables_json",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _safe_int(value: Any) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _clamp_score(value: Any, default: float = DEFAULT_SCORE_THRESHOLD) -> float:
    try:
        score = float(value if value is not None else default)
    except (TypeError, ValueError):
        score = default
    return max(0.0, min(1.0, score))


def _resolve_templates_path(model_path: str | None) -> Path:
    if not model_path:
        return DEFAULT_TEMPLATES_PATH
    candidate = Path(str(model_path).strip())
    if candidate.is_absolute():
        return candidate
    for base in (SCRIPT_DIR, SCRIPT_DIR.parent, SCRIPT_DIR.parent.parent):
        resolved = (base / candidate).resolve()
        if resolved.exists():
            return resolved
    return (SCRIPT_DIR / candidate).resolve()


@lru_cache(maxsize=8)
def _load_templates_cached(path_text: str, mtime_ns: int) -> tuple[TemplateSpec, ...]:
    _ = mtime_ns
    return tuple(load_templates(path_text))


def _get_templates(model_path: str | None) -> tuple[list[TemplateSpec], dict[str, Any]]:
    path = _resolve_templates_path(model_path)
    if not path.is_file():
        raise FileNotFoundError(f"APT模板库文件不存在: {path}")
    stat = path.stat()
    templates = list(_load_templates_cached(str(path), stat.st_mtime_ns))
    meta = {
        "templates_path": str(path),
        "template_count": len(templates),
    }
    return templates, meta


def _normalize_domains(raw_domains: Iterable[Any]) -> tuple[list[str], dict[str, int]]:
    domains: list[str] = []
    seen: set[str] = set()
    raw_count = 0
    invalid_count = 0
    duplicate_count = 0
    for raw in raw_domains:
        raw_count += 1
        normalized = domain_utils.normalize_domain(raw)
        if not domain_utils.is_normalized_domain_like(normalized):
            invalid_count += 1
            continue
        if normalized in seen:
            duplicate_count += 1
            continue
        seen.add(normalized)
        domains.append(normalized)
    return domains, {
        "raw_count": raw_count,
        "valid_count": len(domains),
        "invalid_count": invalid_count,
        "duplicate_count": duplicate_count,
    }


def read_domains_from_file(file_content: bytes, filename: str) -> list[str]:
    file_ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if file_ext == "txt":
        text = file_content.decode("utf-8", errors="ignore")
        raw_domains = [line.strip() for line in text.splitlines() if line.strip()]
    elif file_ext == "csv":
        df = pd.read_csv(io.BytesIO(file_content))
        column = "domain" if "domain" in df.columns else df.columns[0]
        raw_domains = df[column].dropna().astype(str).tolist()
    elif file_ext == "xlsx":
        df = pd.read_excel(io.BytesIO(file_content))
        column = "domain" if "domain" in df.columns else df.columns[0]
        raw_domains = df[column].dropna().astype(str).tolist()
    else:
        raise ValueError(f"不支持的文件类型: {file_ext}")

    domains, _ = _normalize_domains(raw_domains)
    if not domains:
        raise ValueError("文件中没有找到有效域名")
    return domains


def _score_template_match(match: TemplateMatch, template_index: dict[str, TemplateSpec]) -> float:
    spec = template_index.get(match.template)
    if spec is None:
        return DEFAULT_SCORE_THRESHOLD
    score = 0.86
    if spec.suffix_template != "{tld}":
        score += 0.04
    score += min(0.06, spec.fixed_token_count * 0.015)
    score -= min(0.08, spec.wildcard_count * 0.02)
    if spec.wildcard_count == 0:
        score = max(score, 0.98)
    if match.root_count >= 5 or match.row_count >= 10:
        score += 0.02
    return round(_clamp_score(score), 6)


def _risk_level(score: float, matched: bool) -> str:
    if not matched:
        return "low"
    if score >= 0.90:
        return "high"
    if score >= 0.75:
        return "medium"
    return "low"


def _risk_level_label(value: str) -> str:
    return {
        "high": "高",
        "medium": "中",
        "low": "低",
    }.get(value, value)


def _match_reason(match: TemplateMatch) -> str:
    reason = str(match.reason or "").strip()
    if reason:
        return reason
    return f"命中APT注册模板 {match.template}"


def _match_to_row(domain: str, match: TemplateMatch, template_index: dict[str, TemplateSpec]) -> dict[str, Any]:
    parsed = _parse_domain_value(domain)
    registered_domain = f"{parsed.sld}.{parsed.suffix}" if parsed.suffix else parsed.sld
    score = _score_template_match(match, template_index)
    risk_level = _risk_level(score, True)
    reason = _match_reason(match)
    return {
        "域名": domain,
        "规范化域名": domain,
        "注册域名": registered_domain,
        "SLD": parsed.sld,
        "后缀": parsed.suffix,
        "score": score,
        "risk_level": risk_level,
        "风险等级": _risk_level_label(risk_level),
        "reason": reason,
        "命中原因": reason,
        "匹配模板": match.template,
        "模板命中注册域名数": match.root_count,
        "模板命中原始行数": match.row_count,
        "变量JSON": json.dumps(match.variables, ensure_ascii=False, sort_keys=True),
        "预测标签": 1,
        "预测结果": "APT模板命中",
    }


def _normal_row(domain: str) -> dict[str, Any]:
    try:
        parsed = _parse_domain_value(domain)
        registered_domain = f"{parsed.sld}.{parsed.suffix}" if parsed.suffix else parsed.sld
        sld = parsed.sld
        suffix = parsed.suffix
    except Exception:
        registered_domain = domain
        sld = ""
        suffix = ""
    return {
        "域名": domain,
        "规范化域名": domain,
        "注册域名": registered_domain,
        "SLD": sld,
        "后缀": suffix,
        "score": 0.0,
        "risk_level": "low",
        "风险等级": "低",
        "reason": "",
        "命中原因": "",
        "匹配模板": "",
        "模板命中注册域名数": 0,
        "模板命中原始行数": 0,
        "变量JSON": "{}",
        "预测标签": 0,
        "预测结果": "未命中",
    }


def _build_excel(
    *,
    domains: list[str],
    result_rows: list[dict[str, Any]],
    alert_rows: list[dict[str, Any]],
    input_stats: dict[str, int],
    template_meta: dict[str, Any],
    score_threshold: float,
) -> tuple[bytes, dict[str, Any], dict[str, Any]]:
    total = len(domains)
    matched_count = sum(1 for row in result_rows if int(row.get("预测标签") or 0) == 1)
    high_risk_count = len(alert_rows)
    normal_count = max(0, total - matched_count)
    matched_rate = matched_count / total * 100 if total else 0.0
    high_risk_rate = high_risk_count / total * 100 if total else 0.0
    template_counts = Counter(str(row.get("匹配模板") or "") for row in alert_rows if row.get("匹配模板"))

    stats_rows = [
        ("总域名数", total),
        ("APT模板命中域名数", matched_count),
        ("高风险域名数", high_risk_count),
        ("正常域名数", normal_count),
        ("APT模板命中域名占比", f"{matched_rate:.2f}%"),
        ("高风险域名占比", f"{high_risk_rate:.2f}%"),
        ("预警阈值", f"{score_threshold:.2f}"),
        ("模板数量", template_meta.get("template_count", 0)),
        ("无效输入数", input_stats.get("invalid_count", 0)),
        ("重复输入数", input_stats.get("duplicate_count", 0)),
    ]

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(result_rows).to_excel(writer, sheet_name="预测结果", index=False)
        pd.DataFrame(stats_rows, columns=["统计项", "数值"]).to_excel(writer, sheet_name="统计信息", index=False)
        pd.DataFrame(alert_rows).to_excel(writer, sheet_name="APT模板命中域名列表", index=False)

    statistics = {
        "总域名数": total,
        "APT模板命中域名数": matched_count,
        "高风险域名数": high_risk_count,
        "正常域名数": normal_count,
        "APT模板命中域名占比": f"{matched_rate:.2f}%",
        "高风险域名占比": f"{high_risk_rate:.2f}%",
        "total": total,
        "apt_template_nrd": high_risk_count,
        "matched": matched_count,
        "benign": normal_count,
        "score_threshold": score_threshold,
    }
    meta = {
        **template_meta,
        "input_stats": input_stats,
        "matched_count": matched_count,
        "high_risk_count": high_risk_count,
        "template_match_counts": dict(sorted(template_counts.items())),
        "parameters": {
            "score_threshold": score_threshold,
        },
    }
    return output.getvalue(), statistics, meta


def predict_from_domains(
    domains: list[str],
    source_label: str | None = None,
    model_path: str | None = None,
    *,
    score_threshold: float | None = None,
    high_risk_only: bool = False,
) -> tuple[bytes, dict[str, Any], dict[str, Any], list[dict[str, Any]]]:
    normalized_domains, input_stats = _normalize_domains(domains)
    if not normalized_domains:
        raise ValueError("域名列表为空或没有有效域名")

    resolved_threshold = _clamp_score(score_threshold)
    templates, template_meta = _get_templates(model_path)
    template_index = {template.template: template for template in templates}

    all_rows: list[dict[str, Any]] = []
    alert_rows: list[dict[str, Any]] = []
    for domain in normalized_domains:
        match = match_domain(domain, templates)
        if match is None:
            row = _normal_row(domain)
        else:
            row = _match_to_row(domain, match, template_index)
            if float(row.get("score") or 0.0) >= resolved_threshold:
                alert_rows.append(row)
        if not high_risk_only or int(row.get("预测标签") or 0) == 1:
            all_rows.append(row)

    excel_content, statistics, meta = _build_excel(
        domains=normalized_domains,
        result_rows=all_rows,
        alert_rows=alert_rows,
        input_stats=input_stats,
        template_meta=template_meta,
        score_threshold=resolved_threshold,
    )
    meta["source_label"] = source_label or ""
    return excel_content, statistics, meta, alert_rows


def predict_from_file(
    file_content: bytes,
    filename: str,
    model_path: str | None = None,
    *,
    score_threshold: float | None = None,
    high_risk_only: bool = False,
) -> tuple[bytes, dict[str, Any], dict[str, Any], list[dict[str, Any]]]:
    domains = read_domains_from_file(file_content, filename)
    return predict_from_domains(
        domains,
        filename,
        model_path,
        score_threshold=score_threshold,
        high_risk_only=high_risk_only,
    )


def alert_rows_to_score_records(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for row in rows or []:
        domain = str(row.get("域名") or row.get("domain") or "").strip()
        if not domain:
            continue
        try:
            score = float(row.get("score") or 0.0)
        except (TypeError, ValueError):
            score = 0.0
        records.append(
            {
                "domain": domain,
                "score": score,
                "risk_score": score,
                "risk_level": row.get("risk_level") or "",
                "matched_template": row.get("匹配模板") or "",
                "reason": row.get("reason") or row.get("命中原因") or "",
                "raw": json.loads(json.dumps(row, ensure_ascii=False, default=str)),
            }
        )
    return records


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Match NRD domains in a YYYYMMDD-YYYYMMDD range against APT registration templates."
    )
    parser.add_argument("date_range", help="Inclusive date range, for example 20260301-20260302.")
    parser.add_argument("--templates", default=str(DEFAULT_TEMPLATES_PATH), help="Template workbook path.")
    parser.add_argument("--nrd-root", default=str(DEFAULT_NRD_ROOT), help="NRD root directory.")
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT), help="Output root directory.")
    parser.add_argument("--allow-missing", action="store_true", help="Skip missing daily NRD zip files.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    try:
        summary = run_detection(
            date_range=args.date_range,
            templates_path=args.templates,
            nrd_root=args.nrd_root,
            output_root=args.output_root,
            allow_missing=args.allow_missing,
        )
    except Exception as exc:
        print(f"Error: {exc}")
        raise SystemExit(1) from exc

    print(f"Date range: {summary['date_range']}")
    print(f"Templates loaded: {summary['template_count']}")
    print(f"NRD files scanned: {len(summary['nrd_files'])}")
    print(f"NRD domains scanned: {summary['scanned_domains']}")
    print(f"Matched domains: {summary['matched_domains']}")
    print(f"Matches CSV: {summary['output_matches_file']}")
    print(f"Summary JSON: {summary['summary_file']}")


if __name__ == "__main__":
    main()
